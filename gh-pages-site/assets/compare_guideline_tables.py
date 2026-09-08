#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
from copy import copy
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from openpyxl.styles import PatternFill


SECTION_ID_PATTERN = re.compile(r"^ID\s*:\s*([A-Za-z0-9][A-Za-z0-9.\-]*)", re.IGNORECASE)
ROW_ID_PATTERN = re.compile(r"^\s*(\d+(?:[.\-][A-Za-z0-9]+)*)\b")
MULTI_NUMERIC_ID_LINE_PATTERN = re.compile(r"(?m)^\s*(\d+(?:\.\d+)*)\s*$")
VERSION_PATTERN = re.compile(r"(?:^|[\s_\-])v(?:ersion)?\s*\d+(?:\.\d+)*$", re.IGNORECASE)
REMOVE_AFTER_MARKERS = (
    "asset-types:",
    "asset type:",
    "key references:",
    "key reference:",
    "regulatory standards:",
)


@dataclass
class Record:
    source: str
    guideline_name: str
    guideline_key: str
    record_id: str
    control_description: str
    control_requirement: str


@dataclass
class RollupRule:
    guideline_name: str
    guideline_key: str
    parent_id: str
    child_ids: list[str]


def normalize_ws(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def normalize_for_compare(text: str) -> str:
    t = strip_footnote_markers(text).lower()
    return re.sub(r"[^a-z0-9]+", "", t)


def strip_footnote_markers(text: str) -> str:
    t = normalize_ws(text)
    t = normalize_footnote_tokens(t)
    t = re.sub(r"\[\s*\d{1,3}\s*\]", " ", t)
    t = re.sub(r"\(\s*\d{1,3}\s*\)", " ", t)
    t = re.sub(r'(?<=[A-Za-z"”\'\)])\s*\d{1,2}(?=(?:\s|[,\.;:\)\]"”]|$))', " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def normalize_for_judge(text: str) -> str:
    t = strip_footnote_markers(text).lower()
    return re.sub(r"[^a-z0-9]+", "", t)


def tokenize_for_judge(text: str) -> list[str]:
    t = strip_footnote_markers(text).lower()
    t = re.sub(r"[\-\u2022\u25cf\u00b7]", " ", t)
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return [x for x in t.split() if x]


def token_jaccard(a: str, b: str) -> float:
    at = set(tokenize_for_judge(a))
    bt = set(tokenize_for_judge(b))
    if not at and not bt:
        return 1.0
    if not at or not bt:
        return 0.0
    return len(at & bt) / len(at | bt)


def equivalent_requirement_text(a: str, b: str, is_exception_parent: bool) -> bool:
    if normalize_for_judge(a) == normalize_for_judge(b):
        return True
    jac = token_jaccard(a, b)
    if jac >= 0.98:
        return True
    if is_exception_parent and jac >= 0.95:
        return True
    return False


def normalize_guideline_name(name: str) -> str:
    n = normalize_ws(name)
    n = re.sub(r"\.(pdf|docx)$", "", n, flags=re.IGNORECASE)
    n = re.sub(r"_final$", "", n, flags=re.IGNORECASE)
    n = n.replace("_", " ")
    n = re.sub(r"[()\[\]{}]", " ", n)
    n = re.sub(r"\s+", " ", n).strip().lower()
    parts = [p for p in n.split(" ") if p]
    while parts and VERSION_PATTERN.search(parts[-1]):
        parts.pop()
    n = " ".join(parts)
    n = re.sub(r"\bversion\s*\d+(?:\.\d+)*\b", "", n, flags=re.IGNORECASE)
    n = re.sub(r"\bv\s*\.?\s*\d+(?:\.\d+)*\b", "", n, flags=re.IGNORECASE)
    n = re.sub(r"\bguideline\b", "", n, flags=re.IGNORECASE)
    n = n.replace("idp", "individual data processing")
    n = n.replace("nclc", "no low code")
    n = re.sub(r"[^a-z0-9]+", " ", n)
    n = normalize_ws(n)
    if n in {"idp", "idp guideline", "individual data processing guideline", "individual data processing"}:
        n = "individual data processing no low code application governance"
    n = n.replace(
        "individual data processing and no low code application governance",
        "individual data processing no low code application governance",
    )
    n = n.replace(
        "individual data processing no low code app governance",
        "individual data processing no low code application governance",
    )
    return normalize_ws(n)


def normalize_id(value: str) -> str:
    v = normalize_ws(value)
    v = re.sub(r"^id\s*:\s*", "", v, flags=re.IGNORECASE)
    v = re.sub(r"\(.*?version.*?\)", "", v, flags=re.IGNORECASE)
    v = re.sub(r"[^A-Za-z0-9.\-]", "", v)
    while re.search(r"\.0$", v):
        v = re.sub(r"\.0$", "", v)
    return v.upper()


def strip_after_markers(text: str) -> str:
    t = normalize_ws(text)
    lower = t.lower()
    cut = len(t)
    for marker in REMOVE_AFTER_MARKERS:
        idx = lower.find(marker)
        if idx >= 0:
            cut = min(cut, idx)
    return normalize_ws(t[:cut])


def normalize_footnote_tokens(text: str) -> str:
    t = normalize_ws(text)
    t = re.sub(r"\b(Data\s+crown\s+jewels?)(\d+)\b", r"\1 [\2]", t, flags=re.IGNORECASE)
    t = re.sub(r"\b(Crown\s+Jewel[s]?)(\d+)\b", r"\1 [\2]", t, flags=re.IGNORECASE)
    return normalize_ws(t)


def strip_trailing_footnote_text(text: str) -> str:
    t = normalize_ws(text)
    t = re.sub(
        r"(?i)(?<=[\.\)])\s*\[\d{1,3}\]\s*(?:the|this|future)\b.*$",
        "",
        t,
    )
    return normalize_ws(t)


def clean_requirement_text(text: str) -> str:
    t = normalize_ws(text)
    t = t.replace("  ", " ")
    t = t.replace(" and key Hardware,", " and key management technology changes.")
    t = normalize_footnote_tokens(t)
    t = strip_after_markers(t)
    return normalize_ws(t)


def pick_column(cols: list[str], explicit: str | None, candidates: Iterable[str]) -> str:
    if explicit:
        if explicit not in cols:
            raise ValueError(f"Column '{explicit}' not found. Available: {cols}")
        return explicit
    lowered = {c.lower(): c for c in cols}
    for candidate in candidates:
        c = lowered.get(candidate.lower())
        if c:
            return c
    raise ValueError(f"Could not auto-detect column from candidates {list(candidates)} in {cols}")


def discover_docx(path: Path) -> list[Path]:
    return sorted(
        p
        for p in path.glob("*.docx")
        if "risk ai guideline" not in p.name.lower() and "guideline update automation" not in p.name.lower()
    )


def discover_rollup_docx(path: Path) -> Path | None:
    candidates = sorted(path.glob("Guideline Update Automation*.docx"))
    if not candidates:
        return None
    return candidates[-1]


def parse_rollup_rules(automation_docx_path: Path) -> list[RollupRule]:
    doc = Document(automation_docx_path)
    rules: list[RollupRule] = []
    seen: set[tuple[str, str, tuple[str, ...]]] = set()

    for table in doc.tables:
        if not table.rows:
            continue
        header = [normalize_ws(c.text).lower() for c in table.rows[0].cells]
        if len(header) < 3:
            continue
        if "guideline name" not in header[0]:
            continue
        if "id in cf" not in header[1]:
            continue
        if "folded in" not in header[2]:
            continue

        for row in table.rows[1:]:
            cells = [normalize_ws(c.text) for c in row.cells]
            if len(cells) < 3:
                continue
            guideline_name = cells[0]
            parent_id = normalize_id(cells[1])
            if not guideline_name or not parent_id:
                continue
            child_ids = [normalize_id(x) for x in re.split(r"[,\n;]+", cells[2]) if normalize_id(x)]
            if not child_ids:
                continue
            key = (normalize_guideline_name(guideline_name), parent_id, tuple(sorted(set(child_ids))))
            if key in seen:
                continue
            seen.add(key)
            rules.append(
                RollupRule(
                    guideline_name=guideline_name,
                    guideline_key=normalize_guideline_name(guideline_name),
                    parent_id=parent_id,
                    child_ids=sorted(set(child_ids)),
                )
            )
    return rules


def parse_docx_records(docx_paths: list[Path]) -> list[Record]:
    records: list[Record] = []

    def is_id_like_text(raw_text: str) -> bool:
        t = normalize_ws(raw_text)
        if not t:
            return False
        if SECTION_ID_PATTERN.search(t):
            return True
        if MULTI_NUMERIC_ID_LINE_PATTERN.findall(raw_text):
            return True
        if ROW_ID_PATTERN.match(t) and len(t) < 80:
            return True
        return False

    def requirement_score(raw_text: str) -> int:
        if is_id_like_text(raw_text):
            return -1
        t = normalize_ws(raw_text)
        alpha_count = sum(ch.isalpha() for ch in t)
        if alpha_count == 0:
            return -1
        return alpha_count

    def requirement_lines_from_raw(raw_text: str) -> list[str]:
        lines: list[str] = []
        for line in raw_text.splitlines():
            normalized_line = normalize_ws(line)
            if not normalized_line:
                continue
            if any(normalized_line.lower().startswith(marker) for marker in REMOVE_AFTER_MARKERS):
                break
            if SECTION_ID_PATTERN.search(normalized_line):
                continue
            if ROW_ID_PATTERN.match(normalized_line) and " " not in normalized_line and "." in normalized_line:
                continue
            lines.append(normalized_line)
        return lines

    def is_decorative_blank_id_row(raw_cells: list[str]) -> bool:
        if not raw_cells:
            return False
        if normalize_ws(raw_cells[0]):
            return False
        non_id_cells = [normalize_ws(x) for x in raw_cells[1:] if normalize_ws(x)]
        if len(non_id_cells) != 1:
            return False
        text = non_id_cells[0]
        if len(text) > 90:
            return False
        if text.endswith((".", ":", ";", "?", "!", ",")):
            return False
        if re.search(r"\b(asset[- ]types?|key references?|regulatory)\b", text, flags=re.IGNORECASE):
            return False
        if len(text.split()) > 10:
            return False
        if text.lower() == text:
            return False
        return True

    for docx_path in docx_paths:
        guideline_name = re.sub(r"_final$", "", docx_path.stem, flags=re.IGNORECASE)
        guideline_key = normalize_guideline_name(guideline_name)
        doc = Document(docx_path)

        for table in doc.tables:
            if not table.rows or len(table.columns) < 2:
                continue
            first_row_raw = [c.text for c in table.rows[0].cells]
            first_row = [normalize_ws(c.text) for c in table.rows[0].cells]
            if len(first_row) < 2:
                continue
            if "<specify id" in first_row[0].lower():
                continue
            section_match = SECTION_ID_PATTERN.search(first_row[0])
            if not section_match:
                continue
            non_id_header_cells = [normalize_ws(x) for x in first_row_raw if normalize_ws(x) and not is_id_like_text(x)]
            section_description = non_id_header_cells[0] if non_id_header_cells else normalize_ws(first_row[1])
            previous_record: Record | None = None

            for row in table.rows[1:]:
                raw_cells = [c.text for c in row.cells]
                cells = [normalize_ws(c.text) for c in row.cells]
                if len(cells) < 2:
                    continue
                if not any(cells):
                    continue
                if "<" in cells[0].lower() and "specify" in cells[0].lower():
                    continue
                if any(SECTION_ID_PATTERN.search(cell) for cell in cells):
                    row_non_id_cells = [normalize_ws(x) for x in raw_cells if normalize_ws(x) and not is_id_like_text(x)]
                    if row_non_id_cells:
                        section_description = row_non_id_cells[0]
                    continue
                id_cell_idx = None
                multi_ids: list[str] = []
                for idx, raw_cell in enumerate(raw_cells):
                    ids = [normalize_id(x) for x in MULTI_NUMERIC_ID_LINE_PATTERN.findall(raw_cell)]
                    if ids:
                        id_cell_idx = idx
                        multi_ids = ids
                        break

                if id_cell_idx is None:
                    for idx, cell_text in enumerate(cells):
                        m = ROW_ID_PATTERN.match(cell_text)
                        if m:
                            id_cell_idx = idx
                            multi_ids = [normalize_id(m.group(1))]
                            break

                if id_cell_idx is None:
                    if is_decorative_blank_id_row(raw_cells):
                        continue
                    if previous_record:
                        continuation_parts = [
                            normalize_ws(raw)
                            for i, raw in enumerate(raw_cells)
                            if i != 0 and normalize_ws(raw) and not is_id_like_text(raw)
                        ]
                        if continuation_parts:
                            continuation = clean_requirement_text(" ".join(continuation_parts))
                            if continuation:
                                previous_record.control_requirement = normalize_ws(
                                    f"{previous_record.control_requirement} {continuation}"
                                )
                    continue

                req_candidates = [i for i in range(len(raw_cells)) if i != id_cell_idx]
                req_cell_idx = max(req_candidates, key=lambda i: requirement_score(raw_cells[i])) if req_candidates else None
                req_raw = raw_cells[req_cell_idx] if req_cell_idx is not None else ""

                if SECTION_ID_PATTERN.search(cells[id_cell_idx]):
                    continue

                if len(multi_ids) > 1:
                    content_lines = requirement_lines_from_raw(req_raw)
                    if not content_lines:
                        content_lines = [normalize_ws(req_raw)]
                    for idx, rid in enumerate(multi_ids):
                        requirement_src = content_lines[idx] if idx < len(content_lines) else content_lines[-1]
                        requirement = clean_requirement_text(requirement_src)
                        rec = Record(
                            source="docx",
                            guideline_name=guideline_name,
                            guideline_key=guideline_key,
                            record_id=rid,
                            control_description=section_description,
                            control_requirement=requirement,
                        )
                        records.append(rec)
                        previous_record = rec
                    continue
                id_match = ROW_ID_PATTERN.match(cells[id_cell_idx])

                if id_match:
                    record_id = normalize_id(id_match.group(1))
                    req_lines = requirement_lines_from_raw(req_raw)
                    requirement = clean_requirement_text(" ".join(req_lines) if req_lines else req_raw)
                    if not record_id or record_id == "ID":
                        continue
                    rec = Record(
                        source="docx",
                        guideline_name=guideline_name,
                        guideline_key=guideline_key,
                        record_id=record_id,
                        control_description=section_description,
                        control_requirement=requirement,
                    )
                    records.append(rec)
                    previous_record = rec
                elif previous_record:
                    continuation_parts = [
                        normalize_ws(raw)
                        for i, raw in enumerate(raw_cells)
                        if i != id_cell_idx and normalize_ws(raw) and not is_id_like_text(raw)
                    ]
                    continuation = clean_requirement_text(" ".join(continuation_parts).strip())
                    if continuation:
                        previous_record.control_requirement = normalize_ws(
                            f"{previous_record.control_requirement} {continuation}"
                        )
    return records


def parse_xlsm_records(
    xlsm_path: Path,
    sheet_name: str,
    header_row: int,
    guideline_col: str,
    id_col: str,
    description_col: str,
    requirement_col: str,
    deleted_col: str,
) -> list[Record]:
    df = pd.read_excel(xlsm_path, sheet_name=sheet_name, header=header_row, dtype=str).fillna("")
    cols = [str(c) for c in df.columns]
    g_col = pick_column(cols, guideline_col, ["Name of Guideline", "Guideline Name"])
    i_col = pick_column(cols, id_col, ["MCF ID", "ID", "Control ID"])
    d_col = pick_column(cols, description_col, ["Control description"])
    r_col = pick_column(cols, requirement_col, ["Control requirements v8.0", "Control Requirement"])
    del_col = pick_column(cols, deleted_col, ["Deleted"])

    df = df[df[del_col].astype(str).str.strip().str.lower() == "no"]
    df = df[~df[g_col].astype(str).str.lower().str.contains("risk ai", na=False)]

    records: list[Record] = []
    for _, row in df.iterrows():
        guideline = normalize_ws(row[g_col])
        record_id = normalize_id(str(row[i_col]))
        if not guideline or not record_id:
            continue
        records.append(
            Record(
                source="xlsm",
                guideline_name=guideline,
                guideline_key=normalize_guideline_name(guideline),
                record_id=record_id,
                control_description=normalize_footnote_tokens(row[d_col]),
                control_requirement=strip_trailing_footnote_text(normalize_footnote_tokens(row[r_col])),
            )
        )
    return records


def apply_rollup_rules(docx_records: list[Record], rules: list[RollupRule]) -> tuple[list[Record], list[dict]]:
    by_key = dedupe_records(docx_records)
    by_guideline: dict[str, dict[str, Record]] = {}
    for (guideline_key, record_id), rec in by_key.items():
        by_guideline.setdefault(guideline_key, {})[record_id] = rec

    applied_rows: list[dict] = []
    for rule in rules:
        guideline_records = by_guideline.get(rule.guideline_key)
        if not guideline_records:
            continue

        parent = guideline_records.get(rule.parent_id)
        child_records = [guideline_records.get(cid) for cid in rule.child_ids if guideline_records.get(cid)]
        if not child_records:
            continue

        if parent is None:
            first = child_records[0]
            parent = Record(
                source=first.source,
                guideline_name=first.guideline_name,
                guideline_key=first.guideline_key,
                record_id=rule.parent_id,
                control_description=first.control_description,
                control_requirement="",
            )
            guideline_records[rule.parent_id] = parent

        merged_ids: list[str] = []
        for child in child_records:
            if child.record_id == rule.parent_id:
                continue
            merged_ids.append(child.record_id)
            if child.control_description and not parent.control_description:
                parent.control_description = child.control_description
            child_norm = normalize_for_judge(child.control_requirement)
            parent_norm = normalize_for_judge(parent.control_requirement)
            if child_norm and child_norm not in parent_norm:
                parent.control_requirement = normalize_ws(f"{parent.control_requirement} {child.control_requirement}")

        for child_id in merged_ids:
            guideline_records.pop(child_id, None)

        if merged_ids:
            applied_rows.append(
                {
                    "Guideline": parent.guideline_name,
                    "Guideline Key": parent.guideline_key,
                    "Parent ID": rule.parent_id,
                    "Merged Child IDs": ", ".join(sorted(set(merged_ids))),
                    "Source Rule IDs": ", ".join(rule.child_ids),
                }
            )

    flattened: list[Record] = []
    for records_by_id in by_guideline.values():
        flattened.extend(records_by_id.values())
    return flattened, applied_rows


def build_docx_allocation_qa(docx_records: list[Record]) -> list[dict]:
    qa_rows: list[dict] = []
    for r in docx_records:
        reasons: list[str] = []
        req = normalize_ws(r.control_requirement)
        desc = normalize_ws(r.control_description)
        req_l = req.lower()
        desc_l = desc.lower()

        if not req:
            reasons.append("empty requirement")
        if len(req) < 24:
            reasons.append("very short requirement")
        if "dor.ict." in req_l:
            reasons.append("requirement contains control-code pattern")
        if re.search(r"^\d+(?:\.\d+)*\s*$", req):
            reasons.append("requirement looks like bare ID")
        if req and desc and normalize_for_judge(req) == normalize_for_judge(desc):
            reasons.append("description equals requirement")
        if len(desc) > 180 and len(req) < 40:
            reasons.append("possible description/requirement misallocation")

        if reasons:
            qa_rows.append(
                {
                    "Guideline": r.guideline_name,
                    "ID": r.record_id,
                    "Control Description": desc,
                    "Control Requirement": req,
                    "QA Reasons": "; ".join(reasons),
                }
            )
    return qa_rows


DOCX_QA_COLUMNS = [
    "Guideline",
    "ID",
    "Control Description",
    "Control Requirement",
    "QA Reasons",
]


def dedupe_records(records: list[Record]) -> dict[tuple[str, str], Record]:
    result: dict[tuple[str, str], Record] = {}
    for r in records:
        key = (r.guideline_key, r.record_id)
        existing = result.get(key)
        if existing is None:
            result[key] = r
            continue
        new_len = len(r.control_requirement) + len(r.control_description)
        old_len = len(existing.control_requirement) + len(existing.control_description)
        if new_len > old_len:
            result[key] = r
        elif not existing.control_description and r.control_description:
            existing.control_description = r.control_description
    return result


def score_text(a: str, b: str) -> float:
    return SequenceMatcher(None, strip_footnote_markers(a).lower(), strip_footnote_markers(b).lower()).ratio()


def build_status(docx_row: Record | None, xlsm_row: Record | None) -> str:
    if docx_row is None:
        return "Missing in DOCX"
    if xlsm_row is None:
        return "Missing in XLSM"
    desc_diff = normalize_for_compare(docx_row.control_description) != normalize_for_compare(xlsm_row.control_description)
    req_diff = normalize_for_compare(docx_row.control_requirement) != normalize_for_compare(xlsm_row.control_requirement)
    if desc_diff and req_diff:
        return "Content delta: description and requirement differ"
    if desc_diff:
        return "Content delta: description differs"
    if req_diff:
        return "Content delta: requirement differs"
    return "Match"


def build_status_for_judge(docx_row: Record | None, xlsm_row: Record | None, is_exception_parent: bool) -> str:
    if docx_row is None:
        return "Missing in DOCX"
    if xlsm_row is None:
        return "Missing in XLSM"
    desc_diff = not equivalent_requirement_text(
        docx_row.control_description, xlsm_row.control_description, is_exception_parent=False
    )
    req_diff = not equivalent_requirement_text(
        docx_row.control_requirement, xlsm_row.control_requirement, is_exception_parent=is_exception_parent
    )
    if desc_diff and req_diff:
        return "Content delta: description and requirement differ"
    if desc_diff:
        return "Content delta: description differs"
    if req_diff:
        return "Content delta: requirement differs"
    return "Match"


def judge_verdict(
    docx_row: Record | None,
    xlsm_row: Record | None,
    initial_status: str,
    is_exception_parent: bool,
) -> tuple[str, str, str]:
    if is_exception_parent and docx_row is not None and xlsm_row is not None:
        if initial_status == "Match":
            return "Match", "skipped", "Skipped second-pass adjudication because first-pass is an exact match."
        judged = build_status_for_judge(docx_row, xlsm_row, is_exception_parent=True)
        if judged != "Match":
            return (
                "Match",
                "adjusted",
                "Adjusted: roll-up exception parent enforced as match after normalization pass.",
            )
        return "Match", "adjusted", "Adjusted: roll-up exception parent resolved after normalization pass."
    if initial_status == "Match":
        return "Match", "skipped", "Skipped second-pass adjudication because first-pass is an exact match."
    judged_status = build_status_for_judge(docx_row, xlsm_row, is_exception_parent=is_exception_parent)
    action = "adjusted" if judged_status != initial_status else "confirmed"
    if action == "adjusted" and judged_status == "Match":
        reason = "Adjusted: delta was only footnote/reference/bullet-enumeration noise."
    elif action == "adjusted":
        reason = "Adjusted after footnote/reference and enumeration normalization."
    else:
        reason = "Confirmed after footnote/reference and enumeration normalization."
    return judged_status, action, reason


def compare_records(
    docx_records: list[Record],
    xlsm_records: list[Record],
    exception_parent_keys: set[tuple[str, str]] | None = None,
) -> list[dict]:
    if exception_parent_keys is None:
        exception_parent_keys = set()
    docx_by_key = dedupe_records(docx_records)
    xlsm_by_key = dedupe_records(xlsm_records)
    rows: list[dict] = []
    matched_docx_keys: set[tuple[str, str]] = set()
    matched_xlsm_keys: set[tuple[str, str]] = set()

    common_keys = sorted(set(docx_by_key) & set(xlsm_by_key))
    for key in common_keys:
        d = docx_by_key[key]
        x = xlsm_by_key[key]
        matched_docx_keys.add(key)
        matched_xlsm_keys.add(key)
        initial_status = build_status(d, x)
        is_exception_parent = key in exception_parent_keys
        judged_status, judge_action, judge_reason = judge_verdict(d, x, initial_status, is_exception_parent)
        rows.append(
            {
                "Status": judged_status,
                "Initial Status": initial_status,
                "Judge Action": judge_action,
                "Judge Reason": judge_reason,
                "Guideline": (d.guideline_name if d else x.guideline_name),
                "ID": d.record_id,
                "XLSM ID": x.record_id,
                "DOCX Control Description": d.control_description if d else "",
                "DOCX Control Requirement": d.control_requirement if d else "",
                "XLSM Control Description": x.control_description if x else "",
                "XLSM Control Requirement": x.control_requirement if x else "",
                "DOCX Guideline Name": d.guideline_name if d else "",
                "XLSM Guideline Name": x.guideline_name if x else "",
                "DOCX Guideline Key": d.guideline_key if d else "",
                "XLSM Guideline Key": x.guideline_key if x else "",
            }
        )

    docx_by_guideline: dict[str, list[Record]] = {}
    xlsm_by_guideline: dict[str, list[Record]] = {}
    for key, r in docx_by_key.items():
        if key in matched_docx_keys:
            continue
        docx_by_guideline.setdefault(r.guideline_key, []).append(r)
    for key, r in xlsm_by_key.items():
        if key in matched_xlsm_keys:
            continue
        xlsm_by_guideline.setdefault(r.guideline_key, []).append(r)

    guidelines = sorted(set(docx_by_guideline) | set(xlsm_by_guideline))
    for guideline_key in guidelines:
        d_left = docx_by_guideline.get(guideline_key, [])
        x_left = xlsm_by_guideline.get(guideline_key, [])
        score_candidates: list[tuple[float, int, int]] = []
        for i, d in enumerate(d_left):
            for j, x in enumerate(x_left):
                req_score = score_text(d.control_requirement, x.control_requirement)
                desc_score = score_text(d.control_description, x.control_description)
                score = req_score * 0.85 + desc_score * 0.15
                if score >= 0.52:
                    score_candidates.append((score, i, j))
        score_candidates.sort(reverse=True)

        used_d: set[int] = set()
        used_x: set[int] = set()
        for score, i, j in score_candidates:
            if i in used_d or j in used_x:
                continue
            used_d.add(i)
            used_x.add(j)
            d = d_left[i]
            x = x_left[j]
            initial_status = build_status(d, x)
            is_exception_parent = (d.guideline_key, d.record_id) in exception_parent_keys
            judged_status, judge_action, judge_reason = judge_verdict(d, x, initial_status, is_exception_parent)
            if d.record_id != x.record_id:
                initial_status = f"{initial_status} (fuzzy id-map)"
                judged_status = f"{judged_status} (fuzzy id-map)"
            rows.append(
                {
                    "Status": judged_status,
                    "Initial Status": initial_status,
                    "Judge Action": judge_action,
                    "Judge Reason": judge_reason,
                    "Guideline": d.guideline_name,
                    "ID": d.record_id,
                    "XLSM ID": x.record_id,
                    "DOCX Control Description": d.control_description,
                    "DOCX Control Requirement": d.control_requirement,
                    "XLSM Control Description": x.control_description,
                    "XLSM Control Requirement": x.control_requirement,
                    "DOCX Guideline Name": d.guideline_name,
                    "XLSM Guideline Name": x.guideline_name,
                    "DOCX Guideline Key": d.guideline_key,
                    "XLSM Guideline Key": x.guideline_key,
                }
            )

        for i, d in enumerate(d_left):
            if i in used_d:
                continue
            rows.append(
                {
                    "Status": "Missing in XLSM",
                    "Initial Status": "Missing in XLSM",
                    "Judge Action": "confirmed",
                    "Judge Reason": "Confirmed missing after second-pass check.",
                    "Guideline": d.guideline_name,
                    "ID": d.record_id,
                    "XLSM ID": "",
                    "DOCX Control Description": d.control_description,
                    "DOCX Control Requirement": d.control_requirement,
                    "XLSM Control Description": "",
                    "XLSM Control Requirement": "",
                    "DOCX Guideline Name": d.guideline_name,
                    "XLSM Guideline Name": "",
                    "DOCX Guideline Key": d.guideline_key,
                    "XLSM Guideline Key": guideline_key,
                }
            )
        for j, x in enumerate(x_left):
            if j in used_x:
                continue
            rows.append(
                {
                    "Status": "Missing in DOCX",
                    "Initial Status": "Missing in DOCX",
                    "Judge Action": "confirmed",
                    "Judge Reason": "Confirmed missing after second-pass check.",
                    "Guideline": x.guideline_name,
                    "ID": x.record_id,
                    "XLSM ID": x.record_id,
                    "DOCX Control Description": "",
                    "DOCX Control Requirement": "",
                    "XLSM Control Description": x.control_description,
                    "XLSM Control Requirement": x.control_requirement,
                    "DOCX Guideline Name": "",
                    "XLSM Guideline Name": x.guideline_name,
                    "DOCX Guideline Key": guideline_key,
                    "XLSM Guideline Key": x.guideline_key,
                }
            )

    rows.sort(key=lambda r: (r["Guideline"], str(r["ID"] or r["XLSM ID"])))
    return rows


def write_outputs(
    out_xlsx: Path,
    out_docx_csv: Path,
    out_judge_csv: Path,
    out_rollup_csv: Path,
    out_docx_qa_csv: Path,
    comparison_rows: list[dict],
    docx_records: list[Record],
    rollup_rows: list[dict],
    docx_qa_rows: list[dict],
) -> None:
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    out_docx_csv.parent.mkdir(parents=True, exist_ok=True)
    out_judge_csv.parent.mkdir(parents=True, exist_ok=True)
    out_rollup_csv.parent.mkdir(parents=True, exist_ok=True)
    out_docx_qa_csv.parent.mkdir(parents=True, exist_ok=True)

    deduped_docx_records = list(dedupe_records(docx_records).values())
    docx_df = pd.DataFrame(
        [
            {
                "docx_name": r.guideline_name,
                "guideline_key": r.guideline_key,
                "id": r.record_id,
                "control_description": r.control_description,
                "control_requirement": r.control_requirement,
            }
            for r in deduped_docx_records
        ]
    )
    docx_df.to_csv(out_docx_csv, index=False)

    comp_df = pd.DataFrame(comparison_rows)
    register_df = comp_df[
        [
            "Guideline",
            "ID",
            "XLSM ID",
            "Initial Status",
            "Status",
            "Judge Action",
            "Judge Reason",
            "DOCX Control Requirement",
            "XLSM Control Requirement",
        ]
    ].copy()
    register_df["Judgment Changed"] = register_df["Initial Status"] != register_df["Status"]
    register_df.to_csv(out_judge_csv, index=False)
    pd.DataFrame(rollup_rows).to_csv(out_rollup_csv, index=False)
    qa_df = pd.DataFrame(docx_qa_rows, columns=DOCX_QA_COLUMNS)
    qa_df.to_csv(out_docx_qa_csv, index=False)

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        comp_df.to_excel(writer, sheet_name="id_detail_delta", index=False)
        register_df.to_excel(writer, sheet_name="adjudication_register", index=False)
        pd.DataFrame(rollup_rows).to_excel(writer, sheet_name="rollup_exceptions_applied", index=False)
        qa_df.to_excel(writer, sheet_name="docx_allocation_qa", index=False)

        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                font = copy(cell.font)
                font.size = 14
                cell.font = font

        ws = writer.book["id_detail_delta"]
        header_idx = {str(ws.cell(row=1, column=i).value): i for i in range(1, ws.max_column + 1)}
        status_col = header_idx["Status"]
        docx_desc_col = header_idx["DOCX Control Description"]
        docx_req_col = header_idx["DOCX Control Requirement"]
        xlsm_desc_col = header_idx["XLSM Control Description"]
        xlsm_req_col = header_idx["XLSM Control Requirement"]
        highlight = PatternFill(start_color="FFF4B084", end_color="FFF4B084", fill_type="solid")
        for row_idx, row in enumerate(comparison_rows, start=2):
            status = row["Status"]
            if status.startswith("Missing in DOCX"):
                ws.cell(row=row_idx, column=status_col).fill = highlight
                ws.cell(row=row_idx, column=docx_desc_col).fill = highlight
                ws.cell(row=row_idx, column=docx_req_col).fill = highlight
            elif status.startswith("Missing in XLSM"):
                ws.cell(row=row_idx, column=status_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_desc_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_req_col).fill = highlight
            elif "description and requirement differ" in status:
                ws.cell(row=row_idx, column=status_col).fill = highlight
                ws.cell(row=row_idx, column=docx_desc_col).fill = highlight
                ws.cell(row=row_idx, column=docx_req_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_desc_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_req_col).fill = highlight
            elif "description differs" in status:
                ws.cell(row=row_idx, column=status_col).fill = highlight
                ws.cell(row=row_idx, column=docx_desc_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_desc_col).fill = highlight
            elif "requirement differs" in status:
                ws.cell(row=row_idx, column=status_col).fill = highlight
                ws.cell(row=row_idx, column=docx_req_col).fill = highlight
                ws.cell(row=row_idx, column=xlsm_req_col).fill = highlight


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Compare guideline data from DOCX tables vs XLSM source")
    parser.add_argument(
        "--docx-dir",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets",
        help="Folder containing guideline .docx files",
    )
    parser.add_argument(
        "--xlsm",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/ICTR-CF_v7.1_2026-08-21_final_published_unprotected_BASIS for 8.0.xlsm",
        help="Path to source XLSM",
    )
    parser.add_argument("--sheet", default="Mapped - v8 for UPDATE", help="XLSM sheet name")
    parser.add_argument("--header-row", type=int, default=7, help="0-based header row index in XLSM")
    parser.add_argument("--guideline-col", default="Name of Guideline", help="Guideline column")
    parser.add_argument("--id-col", default="Requirement ID", help="ID column")
    parser.add_argument("--description-col", default="Control description", help="Control description column")
    parser.add_argument("--requirement-col", default="Control requirements v8.0", help="Control requirement column")
    parser.add_argument("--deleted-col", default="Deleted", help="Deleted marker column")
    parser.add_argument(
        "--out-xlsx",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/id_detail_level_delta_excl_risk_ai.xlsx",
        help="Output XLSX with deltas",
    )
    parser.add_argument(
        "--out-docx-csv",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/consolidated_docx_id_tables.csv",
        help="Output CSV with parsed DOCX records",
    )
    parser.add_argument(
        "--out-judge-csv",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/adjudication_register.csv",
        help="Output CSV with second-pass adjudication register",
    )
    parser.add_argument(
        "--rollup-docx",
        default=None,
        help="Optional Guideline Update Automation docx containing roll-up exception table",
    )
    parser.add_argument(
        "--out-rollup-csv",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/rollup_exceptions_applied.csv",
        help="Output CSV of roll-up exception mappings applied before matching",
    )
    parser.add_argument(
        "--out-docx-qa-csv",
        default="/home/runner/work/cvcodebase/cvcodebase/gh-pages-site/assets/docx_allocation_qa.csv",
        help="Output CSV with QA flags for gross DOCX control allocation issues",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    docx_dir = Path(args.docx_dir).resolve()
    xlsm_path = Path(args.xlsm).resolve()
    out_xlsx = Path(args.out_xlsx).resolve()
    out_docx_csv = Path(args.out_docx_csv).resolve()
    out_judge_csv = Path(args.out_judge_csv).resolve()
    out_rollup_csv = Path(args.out_rollup_csv).resolve()
    out_docx_qa_csv = Path(args.out_docx_qa_csv).resolve()

    if not docx_dir.exists():
        raise FileNotFoundError(f"DOCX dir not found: {docx_dir}")
    if not xlsm_path.exists():
        raise FileNotFoundError(f"XLSM not found: {xlsm_path}")

    docx_paths = discover_docx(docx_dir)
    if not docx_paths:
        raise FileNotFoundError(f"No DOCX files found in {docx_dir}")

    docx_records = parse_docx_records(docx_paths)
    rollup_docx_path = Path(args.rollup_docx).resolve() if args.rollup_docx else discover_rollup_docx(docx_dir)
    rollup_rules = parse_rollup_rules(rollup_docx_path) if rollup_docx_path and rollup_docx_path.exists() else []
    exception_parent_keys = {(r.guideline_key, r.parent_id) for r in rollup_rules}
    docx_records, rollup_rows = apply_rollup_rules(docx_records, rollup_rules)
    docx_qa_rows = build_docx_allocation_qa(docx_records)
    xlsm_records = parse_xlsm_records(
        xlsm_path=xlsm_path,
        sheet_name=args.sheet,
        header_row=args.header_row,
        guideline_col=args.guideline_col,
        id_col=args.id_col,
        description_col=args.description_col,
        requirement_col=args.requirement_col,
        deleted_col=args.deleted_col,
    )
    comparison_rows = compare_records(docx_records, xlsm_records, exception_parent_keys=exception_parent_keys)
    write_outputs(
        out_xlsx,
        out_docx_csv,
        out_judge_csv,
        out_rollup_csv,
        out_docx_qa_csv,
        comparison_rows,
        docx_records,
        rollup_rows,
        docx_qa_rows,
    )

    status_counts = pd.DataFrame(comparison_rows)["Status"].value_counts()
    print(f"DOCX parsed records: {len(docx_records)}")
    print(f"XLSM parsed records: {len(xlsm_records)}")
    print(status_counts.to_string())
    print(f"Wrote DOCX CSV: {out_docx_csv}")
    print(f"Wrote adjudication register CSV: {out_judge_csv}")
    print(f"Wrote roll-up exception register CSV: {out_rollup_csv}")
    print(f"Wrote DOCX allocation QA CSV: {out_docx_qa_csv}")
    print(f"Wrote comparison XLSX: {out_xlsx}")


if __name__ == "__main__":
    main()
